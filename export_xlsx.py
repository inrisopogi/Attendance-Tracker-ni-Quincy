"""Build the attendance spreadsheet (.xlsx) from students and records."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
LABELS = [("Present", "P"), ("Absent", "A"), ("Late", "L"), ("Excused", "E")]
FILLS = {"P": "D9F2E3", "A": "F8D7D7", "L": "FBEBC8", "E": "E3DEF8"}


def build_workbook(students, records):
    """students: [{'id','name'}]; records: {'YYYY-MM-DD': {studentId: 'P'|'A'|'L'|'E'}}."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    dates = sorted(records)
    first_date_col = 8  # A=ID, B=Name, C..F=counts, G=percent, H onward = dates
    headers = ["ID", "Name"] + [l for l, _ in LABELS] + ["Attendance %"] + dates
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="2F6FED")
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    last_col = get_column_letter(max(first_date_col, first_date_col + len(dates) - 1))
    first_col = get_column_letter(first_date_col)

    for r, s in enumerate(students, start=2):
        rng = f"${first_col}{r}:${last_col}{r}"
        row = [s["id"], s["name"]]
        row += [f'=COUNTIF({rng},"{code}")' for _, code in LABELS]
        row += [f"=IF(SUM(C{r}:F{r})=0,0,(C{r}+E{r}+F{r})/SUM(C{r}:F{r}))"]
        row += [records[d].get(s["id"], "") for d in dates]
        ws.append(row)
        ws.cell(r, 7).number_format = "0.00%"
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT)
            if c >= first_date_col:
                cell.alignment = Alignment(horizontal="center")
                if cell.value in FILLS:
                    cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])

    n = len(students) + 1
    ws.cell(n + 2, 1, "P = Present, A = Absent, L = Late, E = Excused. "
                      "Attendance % counts Present, Late and Excused as attended.").font = Font(name=FONT, italic=True)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 38
    for c in range(3, 8):
        ws.column_dimensions[get_column_letter(c)].width = 11
    for c in range(first_date_col, first_date_col + len(dates)):
        ws.column_dimensions[get_column_letter(c)].width = 12
    ws.freeze_panes = "C2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
