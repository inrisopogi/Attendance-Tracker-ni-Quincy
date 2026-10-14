"""Read students + attendance from one spreadsheet (.xlsx or .csv).

Expected layout (the file produced by "Download spreadsheet" works as-is):
  ID | Name | ...any other columns... | 2026-10-13 | 2026-10-14 | ...
  1  | Dela Cruz, Juan | ... | P | L
Only ID and Name are required. Any column whose header is a date (YYYY-MM-DD or a
real spreadsheet date) is read as attendance: P/A/L/E or Present/Absent/Late/Excused.
"""
import csv
import datetime as dt
import io
import re

ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
WORDS = {"p": "P", "present": "P", "a": "A", "absent": "A",
         "l": "L", "late": "L", "e": "E", "excused": "E"}


def _clean_id(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _as_date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, str) and ISO.match(v.strip()):
        return v.strip()
    return None


def _rows(data, filename):
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data), data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def parse_spreadsheet(data, filename=""):
    rows = _rows(data, filename)
    if not rows:
        raise ValueError("The file is empty.")

    head, id_col, name_col = None, None, None
    for i, row in enumerate(rows[:10]):
        low = [str(c).strip().lower() if c is not None else "" for c in row]
        if "id" in low and "name" in low:
            head, id_col, name_col = i, low.index("id"), low.index("name")
            break
    if head is None:  # no header row: assume column 1 = ID, column 2 = Name
        head, id_col, name_col = -1, 0, 1

    date_cols = {}
    if head >= 0:
        for c, v in enumerate(rows[head]):
            d = _as_date(v)
            if d:
                date_cols[c] = d

    students, records, seen = [], {}, set()
    for row in rows[head + 1:]:
        if len(row) <= name_col:
            continue
        sid, name = _clean_id(row[id_col]), (str(row[name_col]).strip() if row[name_col] is not None else "")
        if not sid or not name or "/" in sid or sid in seen:
            continue
        if sid.lower().startswith("p = present"):  # legend line from our export
            continue
        seen.add(sid)
        students.append({"id": sid, "name": name})
        for c, day in date_cols.items():
            if c < len(row) and row[c] is not None:
                k = WORDS.get(str(row[c]).strip().lower())
                if k:
                    records.setdefault(day, {})[sid] = k

    if not students:
        raise ValueError("No students found. The sheet needs columns named ID and Name.")
    return {"students": students, "records": records}
